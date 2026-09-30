"""Router: export."""
from __future__ import annotations

from typing import Any

from fastapi import APIRouter, Depends, Form, Query, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response

from app.api.deps import (
    common_template_context,
    get_dashboard_service,
    get_export_service,
    get_insights_service,
    get_produccion_repo,
    get_produccion_pita_repo,
    get_nomina_taller_repo,
    get_semanas_repo,
    get_trabajadores_repo,
    get_trabajos_repo,
)
from app.api.helpers import (
    templates,
    _ctx,
    get_catalogos_repo,
    _dev_page_payload,
    _parse_num,
)
from app.core.exceptions import AppError
from app.db.repository import (
    ProduccionRepo,
    ProduccionPitaRepo,
    NominaTallerRepo,
    SemanasRepo,
    TrabajadoresRepo,
    TrabajosRepo,
    CatalogosRepo,
)
from app.services.dashboard_service import DashboardService
from app.services.export_service import ExportService, linea_tiene_gramos

router = APIRouter()

@router.get("/api/export/estadisticas")
async def api_export_estadisticas(
    request: Request,
    vista: str = Query("trabajador"),
    q: str = Query(""),
):
    from app.services.stats_service import StatsService
    from openpyxl import Workbook
    from io import BytesIO
    form_ids = request.query_params.getlist("semana_id")
    selected_ids = [int(x) for x in form_ids if str(x).isdigit()]
    svc = StatsService()
    if vista == "folio":
        rows = svc.stats_folio(q=q, semana_ids=selected_ids or None)
        headers = ["Folio", "Gramos", "Efectivo", "Semanas", "Inicio", "Materiales", "Trabajadores"]
        def line(r):
            return [r["folio"], r["total_gramos"], r["total_efectivo"], r["n_semanas"],
                    r["inicio"], ", ".join(r["materiales"]), " · ".join(r["trabajadores"])]
    else:
        rows = svc.stats_trabajador(q=q, semana_ids=selected_ids or None)
        headers = ["Nombre", "Ubic", "Gramos", "Efectivo", "Semanas", "Folios", "Primera", "Última"]
        def line(r):
            return [r["nombre"], r["ubic"], r["total_gramos"], r["total_efectivo"],
                    r["n_semanas"], r["n_folios"], r["primera"], r["ultima"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Estadisticas"
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for i, r in enumerate(rows, 2):
        for c, v in enumerate(line(r), 1):
            ws.cell(i, c, v)
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="estadisticas.xlsx"'},
    )



@router.get("/api/export/suministro")
async def api_export_suministro(
    semana_id: int | None = None,
    extras: int = Query(2, ge=0, le=5),
    avances: int = Query(4, ge=2, le=10),
    papel: str = Query("letter"),
    tema: str = Query("material"),
    fuente: str = Query("normal"),
    modelo: int = Query(1, ge=0, le=1),
    mat: int = Query(1, ge=0, le=1),
    obs: int = Query(1, ge=0, le=1),
    folio: int = Query(1, ge=0, le=1),
    generico: int = Query(0, ge=0, le=1),
    reps: int = Query(3, ge=1, le=12),
    prod_repo: ProduccionRepo = Depends(get_produccion_repo),
    sem_repo: SemanasRepo = Depends(get_semanas_repo),
):
    """
    Excel de suministro diario (vertical). generico=1: solo ubic+nombre, N reps.
    """
    avances = max(2, min(int(avances or 4), 10))
    from app.services.dia_service import DiaService
    from app.services.export_service import ExportService
    from app.core.calendar_util import encontrar_semana_actual, today
    from fastapi.responses import Response

    st = DiaService().estado_hoy()
    if semana_id is None:
        semanas = sem_repo.listar()
        actual = encontrar_semana_actual(semanas, today())
        semana_id = int(actual["id"]) if actual else (semanas[0]["id"] if semanas else None)
    lineas: list = []
    codigo = ""
    if semana_id:
        lineas = prod_repo.listar_por_semana(int(semana_id), incluir_terminados=False) or []
        sem = next((s for s in sem_repo.listar() if s.get("id") == semana_id), None) or {}
        codigo = sem.get("codigo") or ""
    # Si no hay líneas de semana, usar trabajos del día
    if not lineas:
        lineas = st.get("trabajos") or []
    content, media, filename = ExportService().export_suministro_diario(
        lineas,
        fecha_iso=st.get("fecha") or "",
        label_dia=st.get("label_dia") or "",
        codigo_semana=codigo,
        filas_extra_por_trabajador=extras,
        n_avances=avances,
        papel=papel,
        tema=tema,
        fuente=fuente,
        mostrar_modelo=bool(modelo),
        mostrar_mat=bool(mat),
        mostrar_obs=bool(obs),
        mostrar_folio=bool(folio),
        generico=bool(generico),
        reps_por_trabajador=reps,
    )
    return Response(
        content=content,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )



@router.get("/api/export/produccion/{semana_id}/captura-manual")
async def api_export_captura_manual(
    semana_id: int,
    sem_repo: SemanasRepo = Depends(get_semanas_repo),
    prod_repo: ProduccionRepo = Depends(get_produccion_repo),
    export_svc: ExportService = Depends(get_export_service),
):
    """Nómina en blanco para captura manual: sin precios, total gramos, fila extra por trabajador."""
    semana = sem_repo.obtener(semana_id)
    if not semana:
        from app.core.exceptions import NotFoundError
        raise NotFoundError("Semana no encontrada")
    from app.services.export_service import filtrar_lineas_export, linea_activa

    # Todas las filas de la semana (para contar cuántas se excluyen)
    todas = prod_repo.listar_por_semana(semana_id, incluir_terminados=True)
    excluidas = [r for r in (todas or []) if not linea_activa(r)]
    # Solo activas en el archivo (también sin gramos); nunca terminados/cerrados
    lineas = filtrar_lineas_export(todas or [], solo_con_gramos=False)
    content, media_type, filename = export_svc.export_captura_manual(
        lineas,
        codigo_semana=str(semana.get("codigo") or ""),
        meta_semana={
            "fecha_inicio": semana.get("fecha_inicio"),
            "fecha_fin": semana.get("fecha_fin"),
            "codigo": semana.get("codigo"),
        },
    )
    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Fajos-Incluidas": str(len(lineas)),
            "X-Fajos-Excluidas-Terminados": str(len(excluidas)),
            "X-Fajos-Msg": (
                f"Se excluyeron {len(excluidas)} folio(s) terminado(s)/cerrado(s). "
                f"No aparecen en captura manual ni deben seguir en «sin avance»."
                if excluidas
                else "Sin folios terminados excluidos."
            ),
        },
    )




@router.get("/api/export/trabajadores")
async def api_export_trabajadores(
    activos: int | None = Query(None, description="1=solo activos, 0=solo bajas, omitir=todos"),
    q: str | None = Query(None),
    repo: TrabajadoresRepo = Depends(get_trabajadores_repo),
    export_svc: ExportService = Depends(get_export_service),
):
    """Excel del catálogo de trabajadores."""
    solo_activos = True if activos == 1 else (False if activos == 0 else False)
    # listar: solo_activos True filtra activos; False lista todos
    if activos is None:
        rows = repo.listar(solo_activos=False, busqueda=q)
        flag = None
        stem = "trabajadores_todos"
    elif activos == 1:
        rows = repo.listar(solo_activos=True, busqueda=q)
        flag = True
        stem = "trabajadores_activos"
    else:
        rows = [r for r in repo.listar(solo_activos=False, busqueda=q) if not r.get("activo")]
        flag = False
        stem = "trabajadores_bajas"
    data, media, fname = export_svc.export_trabajadores(rows, filename_stem=stem, solo_activos=flag)
    return Response(
        content=data,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )



@router.get("/api/export/produccion/{semana_id}")
async def api_export(
    semana_id: int,
    fmt: str = Query("xlsx", pattern="^(xlsx|csv|pdf|json)$"),
    formal: int = Query(0, description="1 = ordenar como nómina formal (ubic, nombre)"),
    sem_repo: SemanasRepo = Depends(get_semanas_repo),
    prod_repo: ProduccionRepo = Depends(get_produccion_repo),
    export_svc: ExportService = Depends(get_export_service),
):
    semana = sem_repo.obtener(semana_id)
    if not semana:
        from app.core.exceptions import NotFoundError
        raise NotFoundError("Semana no encontrada")
    # Activas + con gramos; folios terminados fuera de cualquier export
    lineas = prod_repo.listar_por_semana(semana_id, incluir_terminados=False)
    from app.services.export_service import filtrar_lineas_export
    lineas = filtrar_lineas_export(lineas, solo_con_gramos=True)
    if formal:
        def _key(row: dict):
            try:
                u = int(row.get("ubic") or 0)
            except (TypeError, ValueError):
                u = 0
            return (u, str(row.get("nombre") or "").upper(), str(row.get("folio") or "").upper())
        lineas = sorted(lineas, key=_key)
    stem = f"nomina_formal_{semana['codigo']}" if formal else f"nomina_{semana['codigo']}"
    content, media_type, filename = export_svc.export_produccion(
        lineas,
        fmt=fmt,  # type: ignore
        codigo_semana=semana["codigo"],
        meta_semana={"fecha_inicio": semana.get("fecha_inicio"), "fecha_fin": semana.get("fecha_fin"), "codigo": semana.get("codigo")},
        filename_stem=stem,
    )
    if formal and fmt in ("xlsx", "pdf"):
        try:
            from app.services.backup_service import BackupService
            BackupService().guardar_export(
                content,
                codigo_semana=str(semana.get("codigo") or ""),
                extension=f".{fmt}",
                prefijo="formal",
            )
        except Exception:
            pass
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------- Error handler helper (used by main) ----------

def error_response(exc: AppError) -> JSONResponse:
    return JSONResponse(status_code=exc.http_status, content=exc.to_dict())


# ---------- Taller (TLL) y Pita (PIT) ----------


@router.get("/api/export/nomina-taller/{semana_id}")
async def api_export_taller(
    semana_id: int,
    include_resumen: int = Query(1),
    force: int = Query(0, description="1 = permitir sin totales Plata/Pita"),
    repo: NominaTallerRepo = Depends(get_nomina_taller_repo),
    pit: ProduccionPitaRepo = Depends(get_produccion_pita_repo),
    plt: ProduccionRepo = Depends(get_produccion_repo),
    sem_repo: SemanasRepo = Depends(get_semanas_repo),
    export_svc: ExportService = Depends(get_export_service),
):
    """Excel formal Taller (último al engrapar). Incluye RESUMEN de las 3 nóminas + calendario."""
    from fastapi.responses import JSONResponse
    sem = sem_repo.obtener(semana_id)
    if not sem:
        from app.core.exceptions import NotFoundError
        raise NotFoundError("Semana no encontrada")
    todas = list(repo.listar_por_semana(semana_id))
    lineas = [
        r for r in todas
        if float(r.get("total") or 0) > 0
        or float(r.get("sueldo") or 0) > 0
        or float(r.get("extras") or 0) > 0
    ]
    t_plt = plt.totales_semana(semana_id)
    t_pit = pit.totales(semana_id)
    t_tll = repo.totales(semana_id)
    nom_plata = float(t_plt.get("total_efectivo") or 0)
    nom_pita = float(t_pit.get("total") or 0)
    nom_taller = float(t_tll.get("total") or 0)
    # Aviso de totales incompletos se hace en el cliente (preview + confirm).
    # force=1 queda disponible por compatibilidad; el export siempre continúa.
    resumen = None
    if include_resumen:
        resumen = {
            "nomina_taller": nom_taller,
            "nomina_plata": nom_plata,
            "nomina_pita": nom_pita,
        }
    codigo = sem.get("codigo") or str(semana_id)
    meta = {
        "fecha_inicio": sem.get("fecha_inicio"),
        "fecha_fin": sem.get("fecha_fin"),
        "codigo": codigo,
    }
    data, media, fname = export_svc.export_nomina_taller(
        lineas,
        codigo_semana=codigo,
        filename_stem=f"nomina_taller_{codigo}",
        resumen=resumen,
        meta_semana=meta,
    )
    return Response(
        content=data,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )



@router.get("/api/export/nomina-pita/{semana_id}/preview")
async def api_export_pita_preview(
    semana_id: int,
    repo: ProduccionPitaRepo = Depends(get_produccion_pita_repo),
):
    """Cuenta filas totales vs exportables (efectivo > 0)."""
    todas = list(repo.listar_por_semana(semana_id))
    exportables = [r for r in todas if float(r.get("efectivo") or 0) > 0]
    return {
        "ok": True,
        "total": len(todas),
        "exportables": len(exportables),
        "omitidas": len(todas) - len(exportables),
        "regla": "Solo filas con efectivo > 0",
    }



@router.get("/api/export/nomina-taller/{semana_id}/preview")
async def api_export_taller_preview(
    semana_id: int,
    repo: NominaTallerRepo = Depends(get_nomina_taller_repo),
    pit: ProduccionPitaRepo = Depends(get_produccion_pita_repo),
    plt: ProduccionRepo = Depends(get_produccion_repo),
):
    """Cuenta filas exportables y totales de las 3 nóminas (para avisar si faltan)."""
    todas = list(repo.listar_por_semana(semana_id))
    exportables = [
        r for r in todas
        if float(r.get("total") or 0) > 0
        or float(r.get("sueldo") or 0) > 0
        or float(r.get("extras") or 0) > 0
    ]
    t_plt = plt.totales_semana(semana_id)
    t_pit = pit.totales(semana_id)
    t_tll = repo.totales(semana_id)
    nom_plata = float(t_plt.get("total_efectivo") or 0)
    nom_pita = float(t_pit.get("total") or 0)
    nom_taller = float(t_tll.get("total") or 0)
    incompleto = nom_plata <= 0 or nom_pita <= 0
    return {
        "ok": True,
        "total": len(todas),
        "exportables": len(exportables),
        "omitidas": len(todas) - len(exportables),
        "regla": "Solo filas con sueldo/extras/total > 0",
        "nomina_taller": nom_taller,
        "nomina_plata": nom_plata,
        "nomina_pita": nom_pita,
        "gran_total": nom_taller + nom_plata + nom_pita,
        "totales_incompletos": incompleto,
        "aviso": (
            "Faltan totales de Plata y/o Pita en el RESUMEN. "
            "Conviene exportar esas nóminas primero; el GRAN TOTAL saldrá incompleto."
            if incompleto else ""
        ),
    }



@router.get("/api/export/nomina-pita/{semana_id}")
async def api_export_pita(
    semana_id: int,
    repo: ProduccionPitaRepo = Depends(get_produccion_pita_repo),
    sem_repo: SemanasRepo = Depends(get_semanas_repo),
    export_svc: ExportService = Depends(get_export_service),
):
    """Excel formal de nómina Pita (imprimible). Sin bloque de totales generales."""
    sem = sem_repo.obtener(semana_id)
    if not sem:
        from app.core.exceptions import NotFoundError
        raise NotFoundError("Semana no encontrada")
    # Regla fija: no exportar filas con efectivo en 0 (solo lo ya liquidado/llenado)
    todas = list(repo.listar_por_semana(semana_id))
    lineas = [
        r for r in todas
        if float(r.get("efectivo") or 0) > 0
    ]
    codigo = sem.get("codigo") or str(semana_id)
    meta = {
        "fecha_inicio": sem.get("fecha_inicio"),
        "fecha_fin": sem.get("fecha_fin"),
        "codigo": codigo,
    }
    data, media, fname = export_svc.export_nomina_pita(
        lineas, codigo_semana=codigo, filename_stem=f"nomina_pita_{codigo}", resumen=None, meta_semana=meta
    )
    return Response(
        content=data,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )



@router.get("/api/export/resumen/{semana_id}")
async def api_export_resumen(
    semana_id: int,
    repo_tll: NominaTallerRepo = Depends(get_nomina_taller_repo),
    repo_pit: ProduccionPitaRepo = Depends(get_produccion_pita_repo),
    plt: ProduccionRepo = Depends(get_produccion_repo),
    sem_repo: SemanasRepo = Depends(get_semanas_repo),
    export_svc: ExportService = Depends(get_export_service),
):
    """Excel solo con el bloque de totales (Taller + Plata + Pita + Total)."""
    sem = sem_repo.obtener(semana_id)
    if not sem:
        from app.core.exceptions import NotFoundError
        raise NotFoundError("Semana no encontrada")
    t_plt = plt.totales_semana(semana_id)
    t_pit = repo_pit.totales(semana_id)
    t_tll = repo_tll.totales(semana_id)
    codigo = sem.get("codigo") or str(semana_id)
    data, media, fname = export_svc.export_resumen_tres_nominas(
        codigo_semana=codigo,
        nomina_taller=float(t_tll.get("total") or 0),
        nomina_plata=float(t_plt.get("total_efectivo") or 0),
        nomina_pita=float(t_pit.get("total") or 0),
        filename_stem=f"resumen_nominas_{codigo}",
    )
    return Response(
        content=data,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _parse_num(val: str | None, default=None, allow_empty_none: bool = True):
    """Convierte string de formulario a float de forma segura."""
    if val is None:
        return default
    s = str(val).strip().replace(",", ".")
    if s == "":
        return None if allow_empty_none else 0.0
    try:
        return float(s)
    except ValueError as e:
        from app.core.exceptions import ValidationAppError
        raise ValidationAppError(f"Valor numérico inválido: {val!r}") from e



@router.get("/api/export/master")
async def api_export_master(
    export_svc: ExportService = Depends(get_export_service),
):
    """Descarga Excel Master con toda la base operativa (multi-hoja)."""
    content, media_type, filename = export_svc.export_master()
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


