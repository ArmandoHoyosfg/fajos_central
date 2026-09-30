"""Helpers compartidos de exportación (filtros de filas, JSON, página oficio)."""
from __future__ import annotations

from decimal import Decimal
from typing import Any

from openpyxl.styles import Border, Font, PatternFill, Side

def linea_tiene_gramos(row: dict) -> bool:
    """True si la línea tiene al menos un día con gramos > 0 (o total_gramos > 0)."""
    try:
        if float(row.get("total_gramos") or 0) > 0:
            return True
    except (TypeError, ValueError):
        pass
    for col in ("gm_sab", "gm_dom", "gm_lun", "gm_mar", "gm_mie", "gm_jue", "gm_vie"):
        try:
            if float(row.get(col) or 0) > 0:
                return True
        except (TypeError, ValueError):
            continue
    return False


def linea_activa(row: dict) -> bool:
    """False si el folio está terminado/cerrado (no debe salir en exportaciones)."""
    notas = str(row.get("notas") or "").lower()
    if "[terminado]" in notas or "[cerrado]" in notas:
        return False
    if any(x in notas for x in ("folio terminado", "folio cerrado", "trabajo terminado")):
        return False
    try:
        ta = row.get("trabajo_activo")
        if ta is not None and int(ta) == 0:
            return False
    except (TypeError, ValueError):
        pass
    if row.get("es_terminado") in (True, 1, "1", "true", "True", "sí", "si"):
        return False
    if row.get("folio_cerrado") in (True, 1, "1"):
        return False
    # activo del trabajo ligado (alias)
    try:
        if row.get("activo_trabajo") is not None and int(row.get("activo_trabajo")) == 0:
            return False
    except (TypeError, ValueError):
        pass
    return True


def filtrar_lineas_export(lineas: list[dict], *, solo_con_gramos: bool = False) -> list[dict]:
    """Filtro unificado para exportaciones Plata: sin terminados; opcional solo con gramos."""
    out = [r for r in (lineas or []) if linea_activa(r)]
    if solo_con_gramos:
        out = [r for r in out if linea_tiene_gramos(r)]
    return out




def _json_default(obj: Any) -> Any:
    if isinstance(obj, Decimal):
        return float(obj)
    if hasattr(obj, "isoformat"):
        return obj.isoformat()
    raise TypeError(f"No serializable: {type(obj)}")


def print_oficio_landscape(ws) -> None:
    """Tamaño oficio, horizontal, márgenes estrechos (impresión típica)."""
    ws.page_setup.orientation = "landscape"
    try:
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    except Exception:
        pass
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2


# Estilos compartidos (mismos valores que ExportService)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
ALT_FILL = PatternFill("solid", fgColor="E8F4FD")
MAT_FILLS = {
    "AG3": PatternFill("solid", fgColor="D6EAF8"),
    "AG4": PatternFill("solid", fgColor="D5F5E3"),
    "DOL": PatternFill("solid", fgColor="FCF3CF"),
    "DLO": PatternFill("solid", fgColor="FAD7A0"),
}
MAT_DEFAULT = PatternFill("solid", fgColor="E8DAEF")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)
