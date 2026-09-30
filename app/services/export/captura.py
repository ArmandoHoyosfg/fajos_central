"""Mixin de exportación: captura manual."""
from __future__ import annotations

from typing import Any
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.core.exceptions import ExportError, ValidationAppError
from app.core.logging_config import get_logger
from app.services.export.helpers import (
    filtrar_lineas_export,
    linea_activa,
    linea_tiene_gramos,
    print_oficio_landscape,
)

logger = get_logger(__name__)


class CapturaExportMixin:
    def export_captura_manual(
        self,
        lineas: list[dict],
        *,
        codigo_semana: str = "",
        meta_semana: dict | None = None,
    ) -> tuple[bytes, str, str]:
        """
        Nómina en blanco para captura manual (papel).
        - Sin folios terminados/cerrados.
        - Sin fórmulas ni totales de pie.
        - Compacta: fuente 9, márgenes estrechos, una hoja si cabe.
        """
        lineas = filtrar_lineas_export(lineas, solo_con_gramos=False)
        data = self._to_xlsx_captura_manual(
            lineas, codigo_semana, meta_semana=meta_semana
        )
        stem = f"nomina_captura_{codigo_semana or 'semana'}"
        return (
            data,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{stem}.xlsx",
        )

    def _to_xlsx_captura_manual(
        self,
        lineas: list[dict],
        codigo_semana: str,
        *,
        meta_semana: dict | None = None,
    ) -> bytes:
        from collections import OrderedDict
        from datetime import datetime
        from openpyxl.utils import get_column_letter
        from app.utils.fechas_nomina import leyenda_nomina

        generado = datetime.now().strftime("%d/%m/%Y %H:%M")
        meta = meta_semana or {}
        ley = leyenda_nomina(
            meta.get("fecha_inicio"), meta.get("fecha_fin"), codigo=codigo_semana
        )

        def _pk(r):
            try:
                u = int(float(r.get("ubic") or 0))
            except (TypeError, ValueError):
                u = 0
            return (u, str(r.get("nombre") or "").upper(), str(r.get("folio") or "").upper())

        # Defensa extra: solo activas
        data_rows = sorted(
            [r for r in (lineas or []) if linea_activa(r)],
            key=_pk,
        )

        groups: OrderedDict[tuple, list] = OrderedDict()
        for r in data_rows:
            try:
                u = int(float(r.get("ubic") or 0))
            except (TypeError, ValueError):
                u = 0
            key = (str(r.get("nombre") or "").strip().upper(), u)
            groups.setdefault(key, []).append(r)

        wb = Workbook()
        ws = wb.active
        ws.title = "Nomina-Captura"

        name_font = Font(name="Calibri", size=9)
        name_font_b = Font(name="Calibri", size=9, bold=True)
        body_font = Font(name="Calibri", size=9)
        thin = self.THIN

        # --- Encabezado compacto (3 filas) ---
        ws.merge_cells("A1:N1")
        ws["A1"] = "TALLER DE FAJOS CENTRAL — CAPTURA MANUAL (PLATA)"
        ws["A1"].font = Font(name="Calibri", bold=True, size=11, color="1F4E79")
        ws.row_dimensions[1].height = 16

        ws.merge_cells("A2:N2")
        titulo_pago = ley.get("titulo_pago") or "Nómina en blanco — captura"
        ws["A2"] = titulo_pago
        ws["A2"].font = Font(name="Calibri", bold=True, size=10, color="C45C26")
        ws.row_dimensions[2].height = 14

        ws.merge_cells("A3:N3")
        ws["A3"] = (
            f"Semana {codigo_semana or '—'}  ·  "
            f"Periodo: {ley.get('periodo_captura') or '—'}  ·  "
            f"{ley.get('cierre') or 'Cierre: viernes'}  ·  "
            f"{ley.get('pago') or 'Pago: sábado'}  ·  "
            f"Anotar gramos a mano · fila crema = folio extra"
        )
        ws["A3"].font = Font(name="Calibri", size=8, color="555555")
        ws["A3"].alignment = Alignment(wrap_text=False, vertical="center")
        ws.row_dimensions[3].height = 14

        headers = [
            "Nombre", "Ubic", "Folio", "Modelo", "Material", "$/Gr",
            "Sáb", "Dom", "Lun", "Mar", "Mié", "Jue", "Vie", "Total Gr",
        ]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.font = Font(name="Calibri", bold=True, size=8, color="FFFFFF")
            cell.fill = self.HEADER_FILL
            cell.border = thin
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 15

        row_i = 5

        def _write_row(r: dict, *, blank_folio: bool = False) -> None:
            nonlocal row_i
            mat = str(r.get("material") or "").upper() if not blank_folio else ""
            fill = self.MAT_FILLS.get(mat, self.MAT_DEFAULT if mat else None)
            if blank_folio:
                fill = PatternFill("solid", fgColor="FFFDE7")  # crema = espacio extra
            tarifa = r.get("tarifa_gr")
            try:
                tarifa_f = float(tarifa) if tarifa is not None else None
                if tarifa_f == 0:
                    tarifa_f = None
            except (TypeError, ValueError):
                tarifa_f = None
            if blank_folio:
                # Fila extra: solo nombre+ubic; sin folio/modelo/mat ni $/Gr
                vals = [
                    r.get("nombre") or "",
                    r.get("ubic"),
                    "",
                    "",
                    "",
                    "",
                ]
            else:
                vals = [
                    r.get("nombre") or "",
                    r.get("ubic"),
                    r.get("folio") or "",
                    r.get("modelo") or "",
                    r.get("material") or "",
                    tarifa_f if tarifa_f is not None else "",
                ]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(row=row_i, column=col, value=v if v is not None else "")
                cell.border = thin
                cell.font = name_font_b if col == 1 else body_font
                cell.alignment = Alignment(
                    horizontal="left" if col == 1 else "center",
                    vertical="center",
                )
                if fill is not None:
                    cell.fill = fill
            # Días + Total Gr: vacíos (sin fórmulas) para anotar a mano
            for col in range(7, 15):
                cell = ws.cell(row=row_i, column=col, value="")
                cell.border = thin
                if fill is not None:
                    cell.fill = fill
            ws.row_dimensions[row_i].height = 14
            row_i += 1

        if not groups:
            ws.cell(row=row_i, column=1, value="(Sin trabajos activos en esta semana)")
            ws.cell(row=row_i, column=1).font = Font(name="Calibri", size=9, italic=True)
            row_i += 1
        else:
            for (nombre_key, ubic), items in groups.items():
                display_nombre = items[0].get("nombre") or nombre_key
                for item in items:
                    _write_row(item, blank_folio=False)
                # Una fila en blanco por trabajador (otro folio)
                _write_row(
                    {"nombre": display_nombre, "ubic": ubic},
                    blank_folio=True,
                )

        # Anchos compactos (menos espacio muerto)
        widths = {
            "A": 16, "B": 5, "C": 9, "D": 11, "E": 8, "F": 5.5,
            "G": 4.5, "H": 4.5, "I": 4.5, "J": 4.5, "K": 4.5, "L": 4.5, "M": 4.5, "N": 7,
        }
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        # Impresión: horizontal, márgenes estrechos, caber en el menor nº de hojas
        ws.page_setup.orientation = "landscape"
        try:
            ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        except Exception:
            pass
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1  # intenta 1 hoja
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.3
        ws.page_margins.bottom = 0.3
        ws.page_margins.header = 0.15
        ws.page_margins.footer = 0.15
        ws.print_title_rows = "1:4"
        try:
            ws.oddFooter.center.text = f"Captura {codigo_semana} · {generado}"
        except Exception:
            pass

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


